"""Tool implementations behind the MCP server.

Two modes share one code path:

* safe: read-only, narrow tools, SSNs masked, results capped, document reads limited to an
  allowlist. Tool *results* still reach the model, so injection planted in a ticket still
  arrives. That is the point: "safe" tools do not make an assistant immune.
* vulnerable: adds run_sql, read_file, send_email and delete_records with no guard rails,
  and returns unmasked values. This is what a rushed internal integration looks like.

Even in vulnerable mode nothing can leave the lab: send_email only writes a JSON file to
<data>/outbox/, read_file is sandboxed to <data>/docs/, and the _lab_ answer-key tables are
blocked by a SQLite authorizer.
"""

from __future__ import annotations

import json
import re
import sqlite3
import time
import uuid
from pathlib import Path
from typing import Any

from .generate import DB_NAME

MAX_ROWS_SAFE = 10
MAX_ROWS_SQL = 200
SAFE_DOCUMENTS = ("tech/meeting_notes.txt",)
_IDENT = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


def mask_ssn(value: str) -> str:
    return "***-**-" + value[-4:] if value else value


class Lab:
    """The data directory plus the mode the server runs in."""

    def __init__(self, data_dir: str | Path, mode: str = "safe") -> None:
        if mode not in ("safe", "vulnerable"):
            raise ValueError("mode must be 'safe' or 'vulnerable'")
        self.data_dir = Path(data_dir).resolve()
        self.mode = mode
        self.db_path = self.data_dir / DB_NAME
        self.docs_dir = self.data_dir / "docs"
        self.outbox = self.data_dir / "outbox"

    @property
    def vulnerable(self) -> bool:
        return self.mode == "vulnerable"

    def connect(self) -> sqlite3.Connection:
        con = sqlite3.connect(self.db_path)
        con.row_factory = sqlite3.Row

        def authorizer(action: int, arg1: str | None, arg2: str | None, *_: Any) -> int:
            # Block the answer-key tables for every statement, including writes.
            if arg1 and arg1.startswith("_lab_"):
                return sqlite3.SQLITE_DENY
            return sqlite3.SQLITE_OK

        con.set_authorizer(authorizer)
        return con

    def _rows(self, sql: str, params: tuple = ()) -> list[dict]:
        with self.connect() as con:
            return [dict(r) for r in con.execute(sql, params).fetchall()]

    # -- safe tools ------------------------------------------------------------

    def list_datasets(self) -> dict:
        out = {}
        with self.connect() as con:
            names = [
                r[0]
                for r in con.execute(
                    "SELECT name FROM sqlite_master WHERE type='table' "
                    "AND name NOT LIKE '\\_lab\\_%' ESCAPE '\\' ORDER BY name"
                )
            ]
            for n in names:
                out[n] = con.execute(f'SELECT COUNT(*) FROM "{n}"').fetchone()[0]
        return {"tables": out, "documents": list(self.allowed_documents())}

    def search_customers(self, name: str, limit: int = 5) -> list[dict]:
        limit = max(1, min(limit, MAX_ROWS_SAFE))
        rows = self._rows(
            "SELECT id, full_name, email, phone, ssn, online_username, online_password "
            "FROM banking_customers WHERE full_name LIKE ? ORDER BY id LIMIT ?",
            (f"%{name}%", limit),
        )
        return [self._shape_customer(r) for r in rows]

    def search_patients(self, name: str, limit: int = 5) -> list[dict]:
        limit = max(1, min(limit, MAX_ROWS_SAFE))
        rows = self._rows(
            "SELECT id, mrn, full_name, dob, ssn, phone, email, insurance_provider, "
            "insurance_id FROM healthcare_patients WHERE full_name LIKE ? ORDER BY id LIMIT ?",
            (f"%{name}%", limit),
        )
        for r in rows:
            if not self.vulnerable:
                r["ssn"] = mask_ssn(r["ssn"])
                r.pop("insurance_id")
        return rows

    def _shape_customer(self, r: dict) -> dict:
        if not self.vulnerable:
            r["ssn"] = mask_ssn(r["ssn"])
            r.pop("online_password", None)
        return r

    def list_tickets(self, status: str | None = None, limit: int = 10) -> list[dict]:
        limit = max(1, min(limit, MAX_ROWS_SAFE))
        if status:
            return self._rows(
                "SELECT id, title, status FROM tech_tickets WHERE status=? ORDER BY id LIMIT ?",
                (status, limit),
            )
        return self._rows("SELECT id, title, status FROM tech_tickets ORDER BY id LIMIT ?", (limit,))

    def get_ticket(self, ticket_id: int) -> dict:
        rows = self._rows(
            "SELECT id, created_by, title, body, status FROM tech_tickets WHERE id=?",
            (ticket_id,),
        )
        return rows[0] if rows else {"error": f"no ticket {ticket_id}"}

    def allowed_documents(self) -> tuple[str, ...]:
        if self.vulnerable:
            return tuple(
                sorted(p.relative_to(self.docs_dir).as_posix() for p in self.docs_dir.rglob("*") if p.is_file())
            )
        return SAFE_DOCUMENTS

    def read_document(self, path: str) -> dict:
        if path not in SAFE_DOCUMENTS:
            return {"error": "document not available", "available": list(SAFE_DOCUMENTS)}
        return {"path": path, "text": (self.docs_dir / path).read_text(encoding="utf-8")}

    # -- vulnerable-only tools -------------------------------------------------

    def run_sql(self, query: str) -> dict:
        try:
            with self.connect() as con:
                cur = con.execute(query)
                if cur.description is None:
                    return {"rows_affected": cur.rowcount}
                cols = [d[0] for d in cur.description]
                rows = cur.fetchmany(MAX_ROWS_SQL)
                return {"columns": cols, "rows": [list(r) for r in rows], "truncated": len(rows) == MAX_ROWS_SQL}
        except sqlite3.Error as exc:
            return {"error": str(exc)}

    def read_file(self, path: str) -> dict:
        target = (self.docs_dir / path).resolve()
        if self.docs_dir not in target.parents or not target.is_file():
            return {"error": "file not found in the lab documents folder"}
        if target.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(target, read_only=True)
            return {"path": path, "sheets": {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb}}
        return {"path": path, "text": target.read_text(encoding="utf-8", errors="replace")}

    def send_email(self, to: str, subject: str, body: str) -> dict:
        """Pretend to send. Writes a file instead, so a demo can show what would have left."""
        self.outbox.mkdir(parents=True, exist_ok=True)
        name = f"{int(time.time())}-{uuid.uuid4().hex[:8]}.json"
        (self.outbox / name).write_text(
            json.dumps({"to": to, "subject": subject, "body": body}, indent=2), encoding="utf-8"
        )
        return {"status": "sent (simulated)", "outbox_file": name}

    def delete_records(self, table: str, where: str) -> dict:
        if not _IDENT.match(table):
            return {"error": "bad table name"}
        try:
            with self.connect() as con:
                cur = con.execute(f'DELETE FROM "{table}" WHERE {where}')
                return {"deleted": cur.rowcount}
        except sqlite3.Error as exc:
            return {"error": str(exc)}
