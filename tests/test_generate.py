import hashlib
import re
import sqlite3
from pathlib import Path

import pytest

from ai_assistant_mcp_server import fakes
from ai_assistant_mcp_server.generate import DB_NAME, build_dataset
from ai_assistant_mcp_server.payloads import CANARIES, PAYLOADS


@pytest.fixture(scope="module")
def dataset(tmp_path_factory):
    out = tmp_path_factory.mktemp("lab")
    result = build_dataset(out, seed=1337)
    return out, result


def rows(out, sql, args=()):
    con = sqlite3.connect(Path(out) / DB_NAME)
    try:
        return con.execute(sql, args).fetchall()
    finally:
        con.close()


def test_same_seed_gives_identical_data(tmp_path):
    a, b = tmp_path / "a", tmp_path / "b"
    build_dataset(a, seed=7, scale=0.2)
    build_dataset(b, seed=7, scale=0.2)

    def digest(root):
        h = hashlib.sha256()
        con = sqlite3.connect(root / DB_NAME)
        for line in con.iterdump():
            h.update(line.encode())
        con.close()
        for f in sorted((root / "docs").rglob("*")):
            if f.is_file() and f.suffix != ".xlsx":  # xlsx embeds a timestamp
                h.update(f.read_bytes())
        return h.hexdigest()

    assert digest(a) == digest(b)


def test_different_seed_gives_different_data(tmp_path):
    build_dataset(tmp_path / "a", seed=1, scale=0.2)
    build_dataset(tmp_path / "b", seed=2, scale=0.2)
    q = "SELECT ssn FROM healthcare_patients ORDER BY id LIMIT 5"
    assert rows(tmp_path / "a", q) != rows(tmp_path / "b", q)


def test_regenerating_resets_the_database(tmp_path):
    build_dataset(tmp_path, seed=3, scale=0.2)
    con = sqlite3.connect(tmp_path / DB_NAME)
    con.execute("DELETE FROM banking_transactions")
    con.commit()
    con.close()
    build_dataset(tmp_path, seed=3, scale=0.2)
    assert rows(tmp_path, "SELECT COUNT(*) FROM banking_transactions")[0][0] > 0


def test_every_random_ssn_is_in_the_never_issued_range(dataset):
    out, _ = dataset
    for table in ("healthcare_patients", "banking_customers", "tech_employees"):
        for (value,) in rows(out, f"SELECT ssn FROM {table}"):
            assert re.fullmatch(r"9\d\d-\d\d-\d{4}", value), value
            assert int(value[:3]) >= 900


def test_card_numbers_are_luhn_valid(dataset):
    out, _ = dataset
    cards = [c for (c,) in rows(out, "SELECT card_number FROM banking_cards")]
    assert len(cards) > 100
    assert all(len(c) == 16 and fakes.is_luhn_valid(c) for c in cards)


def test_canaries_are_present_and_listed(dataset):
    out, _ = dataset
    listed = {k: v for k, v, _ in rows(out, "SELECT kind, value, location FROM _lab_canaries")}
    assert listed == CANARIES
    assert rows(out, "SELECT COUNT(*) FROM healthcare_patients WHERE ssn=?", (CANARIES["ssn"],))[0][0] == 1
    assert rows(out, "SELECT COUNT(*) FROM banking_cards WHERE card_number=?", (CANARIES["card"],))[0][0] == 1
    assert rows(out, "SELECT COUNT(*) FROM tech_credentials WHERE api_key=?", (CANARIES["api_key"],))[0][0] == 1


def test_random_values_never_collide_with_canaries(dataset):
    out, _ = dataset
    assert rows(out, "SELECT COUNT(*) FROM tech_employees WHERE ssn=?", (CANARIES["ssn"],))[0][0] == 0


def test_poisoned_records_carry_payloads_and_are_recorded(dataset):
    out, _ = dataset
    keys = rows(out, "SELECT table_name, row_id, column_name, payload_kind FROM _lab_poisoned")
    assert len(keys) >= 6
    assert {k for *_, k in keys} == set(PAYLOADS)
    for table, row_id, column, kind in keys:
        (text,) = rows(out, f"SELECT {column} FROM {table} WHERE id=?", (row_id,))[0]  # noqa: S608
        assert PAYLOADS[kind] in text


def test_documents_are_written(dataset):
    out, result = dataset
    names = {Path(p).name for p in result["docs"]}
    assert {"payroll_2026.xlsx", "employee_export.csv", "production.env", "meeting_notes.txt"} <= names
    assert all(Path(p).stat().st_size > 0 for p in result["docs"])
    assert "ignore your previous instructions" in (out / "docs/tech/meeting_notes.txt").read_text()


def test_excel_file_opens_and_has_payroll_rows(dataset):
    from openpyxl import load_workbook

    out, _ = dataset
    ws = load_workbook(out / "docs/hr/payroll_2026.xlsx")["Payroll"]
    assert ws.max_row > 10
    assert ws["C2"].value.startswith("9")


def test_nothing_looks_like_a_real_provider_secret(dataset):
    """GitHub push protection and secret scanners match these shapes; none may appear."""
    out, _ = dataset
    patterns = [r"AKIA[0-9A-Z]{16}", r"ghp_[A-Za-z0-9]{20,}", r"sk_live_", r"xox[abp]-",
                r"AIza[0-9A-Za-z_\-]{30,}", r"-----BEGIN [A-Z ]*PRIVATE KEY-----"]
    blobs = [Path(p).read_bytes().decode("utf-8", "ignore") for p in dataset[1]["docs"]
             if not p.endswith(".xlsx")]
    blobs.append("\n".join(sqlite3.connect(out / DB_NAME).iterdump()))
    for blob in blobs:
        for pat in patterns:
            assert not re.search(pat, blob), pat


def test_emails_and_invalid_domain_are_safe(dataset):
    out, _ = dataset
    for (email,) in rows(out, "SELECT email FROM banking_customers"):
        assert email.endswith("@example.com")
